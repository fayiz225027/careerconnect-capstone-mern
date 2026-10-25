import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def md_to_docx(md_path, docx_path, title_text, subtitle_text):
    print(f"Converting {md_path} -> {docx_path}...")
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 41, 59)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_rows = []

    for line in lines:
        line_str = line.strip()

        # Handle Tables
        if line_str.startswith('|') and line_str.endswith('|'):
            if '---' in line_str:
                continue # delimiter line
            cells = [c.strip() for c in line_str.strip('|').split('|')]
            table_rows.append(cells)
            in_table = True
            continue
        else:
            if in_table and table_rows:
                # Flush table to document
                cols_count = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=cols_count)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < cols_count:
                            cell = t.cell(r_idx, c_idx)
                            cell.text = cell_text
                            # Bold header row
                            if r_idx == 0:
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.bold = True
                doc.add_paragraph() # spacing
                table_rows = []
                in_table = False

        if not line_str:
            continue

        # Headings
        if line_str.startswith('# '):
            p = doc.add_heading(level=1)
            run = p.add_run(line_str[2:])
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(2, 132, 199) # Primary Blue
            run.bold = True
        elif line_str.startswith('## '):
            p = doc.add_heading(level=2)
            run = p.add_run(line_str[3:])
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate
            run.bold = True
        elif line_str.startswith('### '):
            p = doc.add_heading(level=3)
            run = p.add_run(line_str[4:])
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(7, 89, 133)
            run.bold = True
        elif line_str.startswith('- ') or line_str.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = line_str[2:]
            # parse bold inside text
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line_str):
            p = doc.add_paragraph(style='List Number')
            num_prefix = re.match(r'^\d+\.\s', line_str).group(0)
            text = line_str[len(num_prefix):]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line_str)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)

    # Flush any remaining table
    if in_table and table_rows:
        cols_count = max(len(r) for r in table_rows)
        t = doc.add_table(rows=len(table_rows), cols=cols_count)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell_text in enumerate(row):
                if c_idx < cols_count:
                    cell = t.cell(r_idx, c_idx)
                    cell.text = cell_text
                    if r_idx == 0:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.bold = True
        table_rows = []

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    base_dir = r'C:\Users\DELL\Desktop\capstone'
    srs_md = os.path.join(base_dir, 'docs', 'SRS_CareerConnect.md')
    srs_docx = os.path.join(base_dir, 'docs', 'SRS_CareerConnect.docx')
    md_to_docx(srs_md, srs_docx, 'Software Requirements Specification', 'CareerConnect MERN Portal')

    report_md = os.path.join(base_dir, 'docs', 'CareerConnect_Project_Report_ICTAK.md')
    report_docx = os.path.join(base_dir, 'docs', 'CareerConnect_Project_Report_ICTAK.docx')
    md_to_docx(report_md, report_docx, 'Summer Internship Report', 'ICT Academy of Kerala')
